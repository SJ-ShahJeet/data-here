"""The prospecting steps. Each one uses brightdata (data), llm (judgment), graph (storage)."""

import argparse
import csv
import json
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from urllib.parse import urljoin, urlparse

from pydantic import BaseModel
from selectolax.lexbor import LexborHTMLParser as HTMLParser

from datahere import luma, brightdata, graph, llm
from datahere.schema import Attendee, CompanyInfo, Match, Note, PersonProfile, SellerProfile, Signal

HEADERS = {
    "name": ("name", "full name", "attendee"),
    "role": ("role", "title", "job title", "position"),
    "company": ("company", "organization", "organisation", "company name"),
}


def match_headers(header: list) -> dict[str, int]:
    """Map name/role/company -> column index, case-insensitive. Name and Company are required."""
    norm = [str(h or "").strip().lower() for h in header]
    cols = {}
    for field, aliases in HEADERS.items():
        for alias in aliases:
            if alias in norm:
                cols[field] = norm.index(alias)
                break
    missing = {"name", "company"} - cols.keys()
    if missing:
        raise ValueError(f"sheet header {header!r} is missing column(s): {sorted(missing)}")
    return cols


def load_sheet(path: Path) -> list[Attendee]:
    """.xlsx or .csv with Name / Role / Company columns (case-insensitive)."""
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
        wb.close()
    else:
        with path.open(newline="", encoding="utf-8-sig") as f:
            rows = list(csv.reader(f))
    if not rows:
        return []
    cols = match_headers(list(rows[0]))

    def cell(row, field):
        i = cols.get(field)
        return str(row[i] or "").strip() if i is not None and i < len(row) else ""

    out = []
    for row in rows[1:]:
        name = cell(row, "name")
        if name:
            out.append(Attendee(name=name, role=cell(row, "role"), company=cell(row, "company")))
    return out


def _clip(s: str, n: int = 300) -> str:
    s = " ".join(str(s or "").split())
    return s if len(s) <= n else s[: n - 1] + "…"


SELLER_SYSTEM = """You profile a B2B company from its website so a sales team can find buyers at an event.
Rules:
- name: the company's short brand name.
- summary: 1-2 sentences, what they sell and to whom.
- icp.segments: lowercase short segment names of the customers they sell to (e.g. "fast casual", "qsr", "casual dining").
- icp.min_locations / max_locations: only if the site implies a customer size, else null.
- icp.buyer_roles: lowercase role keywords of likely buyers (e.g. "cfo", "vp operations", "cmo", "founder", "ceo").
- icp.pain_topics: lowercase 1-2 word keywords for the problems they solve (e.g. "pricing", "margins", "food cost", "promotions").
- icp.disqualify: lowercase keywords for companies that are NOT buyers (e.g. "vendor", "single location").
- customers: real customer brands, usually in "Image labels (logos)" or testimonials. Exclude the seller itself, investors/partners/accelerators, and made-up names shown inside product UI mockups or demo dashboards.
Use only what the pages say."""


def _internal_links(home: str, html: str, limit: int = 3) -> list[str]:
    host = urlparse(home).netloc.removeprefix("www.")
    keys = ("about", "solution", "product", "platform", "how-it-works", "customers", "why")
    seen = []
    for a in HTMLParser(html).css("a[href]"):
        url = urljoin(home, a.attributes.get("href") or "").split("#")[0].rstrip("/")
        if urlparse(url).netloc.removeprefix("www.") != host or url in seen or url == home.rstrip("/"):
            continue
        if any(k in urlparse(url).path.lower() for k in keys):
            seen.append(url)
    return seen[:limit]


def profile_seller(url: str) -> SellerProfile:
    """Scrape the seller's site (home + a few internal pages), LLM -> SellerProfile."""
    if not url.startswith("http"):
        url = "https://" + url
    pages = [f"# {url}\n{brightdata.page_text(url)}"]
    try:
        links = _internal_links(url, brightdata.fetch(url))
    except Exception as e:  # links are a bonus; homepage alone still works
        print(f"  ! link discovery failed: {e}")
        links = []
    for link in links:
        try:
            pages.append(f"# {link}\n{brightdata.page_text(link, max_chars=4000)}")
        except Exception as e:
            print(f"  ! skip {link}: {e}")
    prof = llm.structured("\n\n".join(pages), SellerProfile, system=SELLER_SYSTEM)
    return prof.model_copy(update={"url": url})


def _company_system(seller: SellerProfile) -> str:
    return f"""You triage a company that attended an industry event, for the seller below.
Seller: {seller.name} - {seller.summary}
Seller ICP segments: {seller.icp.segments}
Seller pain topics: {seller.icp.pain_topics}
Rules:
- kind: "operator" ONLY if the company itself runs the business the seller sells into (e.g. it owns or franchises restaurants that serve guests).
  "vendor" if it sells software, platforms, displays, POS, delivery, data, equipment, food supply, consulting or any service TO such businesses,
  even when its site talks about "restaurant operators" or names restaurant clients. "investor" for funds/PE/VC; else association/media/other.
  When unsure between operator and vendor, choose vendor.
- segment: pick one of the seller's ICP segments when it fits, else a short lowercase segment or "".
- locations: number of locations only if stated, else null.
- summary: one sentence. If no website was available, say the classification is from the name only.
- signals: only kind "site", only when the page shows evidence of one of the seller's pain topics; topic = that pain topic, text = a short quote from the page. Empty list if none.
- domain: leave "" (filled by the caller)."""


def triage_company(name: str, seller: SellerProfile) -> CompanyInfo:
    """Find domain, scrape homepage, LLM -> kind/segment/locations/site signals."""
    domain = brightdata.find_company_domain(name)
    page = ""
    if domain:
        try:
            page = brightdata.page_text(domain if domain.startswith("http") else f"https://{domain}", max_chars=5000)
        except Exception as e:
            print(f"  ! {name}: homepage fetch failed: {e}")
    snippets = "\n".join(f"- {r['title']}: {r['snippet']}" for r in brightdata.search(f"{name} official site")[:4])
    prompt = (f"Company: {name}\nDomain: {domain or 'not found'}\n\nSearch results:\n{snippets}\n\n"
              f"Homepage text:\n{page or '(homepage unavailable; use the search results)'}")
    info = llm.structured(prompt, CompanyInfo, system=_company_system(seller))
    signals = [s.model_copy(update={"kind": "site", "text": _clip(s.text), "url": s.url or domain}) for s in info.signals]
    return info.model_copy(update={"name": name, "domain": domain, "signals": signals})


def _first(d: dict, *keys, default=""):
    for k in keys:
        v = d.get(k)
        if v:
            return v
    return default


class _PostTag(BaseModel):
    index: int
    topic: str


class _PostTags(BaseModel):
    tags: list[_PostTag]


def enrich_person(a: Attendee, seller: SellerProfile) -> PersonProfile:
    """Find LinkedIn, pull profile + recent posts, LLM tags posts with seller pain topics."""
    base = PersonProfile(name=a.name, company=a.company, title=a.role)
    url = a.linkedin_url or brightdata.find_linkedin_url(a.name, a.company)
    if not url:
        return base
    prof = (brightdata.linkedin_profiles([url]) or [{}])[0]
    cur = prof.get("current_company") if isinstance(prof.get("current_company"), dict) else {}
    base.company = a.company or _first(cur, "name") or _first(prof, "current_company_name")
    headline = _first(prof, "headline", "position", "about")
    # The profile record already carries recent posts (articles) and activity (shared posts),
    # so one profile call covers intent; no separate posts-dataset call.
    posts = (prof.get("posts") or []) + (prof.get("activity") or [])
    texts = [
        (_clip(" - ".join(x for x in (_first(p, "title"), _first(p, "attribution")) if x)), _first(p, "link", "url"), str(_first(p, "created_at", default=""))[:10])
        for p in posts
    ]
    texts = [t for t in texts if t[0]]
    signals = []
    if texts and seller.icp.pain_topics:
        listing = "\n".join(f"[{i}] {t}" for i, (t, _, _) in enumerate(texts))
        tags = llm.structured(
            f"Topics: {seller.icp.pain_topics}\n\nPosts:\n{listing}",
            _PostTags,
            system=(f"Seller: {seller.name}, {seller.summary}\n"
                    "Tag each post with ONE topic from the list only if the post shows the author or their company "
                    "dealing with that problem in a way this seller could help with (e.g. raising menu prices, margin "
                    "pressure, discounting/promo strategy, food cost increases). Brand launches, hiring, awards, generic "
                    "marketing or company news are 'other'. When unsure, 'other'. Return one tag per post index."),
        )
        topics = {t.lower(): t for t in seller.icp.pain_topics}
        for t in tags.tags:
            topic = topics.get(t.topic.strip().lower())
            if topic and 0 <= t.index < len(texts):
                text, purl, date = texts[t.index]
                signals.append(Signal(kind="post", topic=topic, text=text, url=purl, date=date))
    return base.model_copy(
        update={
            # sheet titles are trusted; Luma 'role' is just a bio, so prefer LinkedIn there
            "title": (a.role if a.company else "") or _first(cur, "title") or _first(prof, "position") or a.role,
            "linkedin_url": _first(prof, "url", "input_url", default=url),
            "headline": _clip(headline, 200),
            "signals": signals,
        }
    )


def parse_note(text: str) -> Note:
    return llm.structured(
        text,
        Note,
        system="Extract a networking note. person = who was met; company if stated; met_at = where/when; "
        "can_help_with = short lowercase topics they can help with; knows = names of other people they know. Use only what is written.",
    )


class _Email(BaseModel):
    subject: str
    body: str


def draft_email(m: Match, seller: SellerProfile) -> dict:
    """Returns {"subject", "body"}; body opens with m.hook when present."""
    reason = m.reasons[0].text if m.reasons else ""
    prompt = f"""Write a short cold email (max 120 words in the body).
To: {m.name}, {m.title} at {m.company}
Opening hook (open the body with this, in your own words): {m.hook or '(none; open with the reason)'}
One reason they are a fit: {reason}
Seller: {seller.name} - {seller.summary}
Seller customers you may name: {seller.customers}
Rules: cite only that one reason, make no claims beyond the seller description, one soft call to action, sign off with "{seller.name} team"."""
    e = llm.structured(prompt, _Email, system="You write brief, specific, honest B2B outreach emails.")
    return {"subject": e.subject, "body": e.body}


def run(seller_url: str, sheet: Path | str, event: str, deep: int = 30) -> dict:
    """Full batch: profile seller, triage all companies, enrich people at the top `deep` operator companies, write graph. Returns stats.
    `sheet` may also be a public Luma event URL (hosts + featured guests only)."""
    graph.init()
    attendees = luma.load(str(sheet)) if str(sheet).startswith("http") else load_sheet(Path(sheet))
    print(f"loaded {len(attendees)} attendees")
    seller = profile_seller(seller_url)
    graph.upsert_seller(seller)
    graph.upsert_event(event)
    print(f"seller: {seller.name} | segments={seller.icp.segments} | topics={seller.icp.pain_topics}")

    names = list({a.company.lower(): a.company for a in attendees if a.company}.values())

    def triage(name):
        try:
            return triage_company(name, seller)
        except Exception as e:
            print(f"  ! triage {name}: {e}")
            return CompanyInfo(name=name, summary=f"triage failed: {e}")

    companies = {}
    with ThreadPoolExecutor(max_workers=8) as pool:
        for i, c in enumerate(pool.map(triage, names), 1):
            graph.upsert_company(c)
            companies[c.name.lower()] = c
            print(f"[{i}/{len(names)}] {c.name}: {c.kind} {c.segment or '-'} {c.locations or ''}")

    segments = {s.lower() for s in seller.icp.segments}
    operators = [c for c in companies.values() if c.kind == "operator"]
    operators.sort(key=lambda c: (c.segment.lower() in segments, c.locations or 0), reverse=True)
    top = {c.name.lower() for c in operators[:deep]}
    to_enrich = [a for a in attendees if a.company.lower() in top]

    def enrich(a):
        try:
            return enrich_person(a, seller)
        except Exception as e:
            print(f"  ! enrich {a.name}: {e}")
            return PersonProfile(name=a.name, company=a.company, title=a.role)

    with ThreadPoolExecutor(max_workers=8) as pool:
        for i, p in enumerate(pool.map(enrich, to_enrich), 1):
            graph.upsert_person(p, event)
            print(f"[{i}/{len(to_enrich)}] enriched {p.name} ({len(p.signals)} signals)")

    for a in attendees:
        if a.company.lower() not in top:
            graph.upsert_person(PersonProfile(name=a.name, company=a.company, title=a.role), event)

    return {**graph.stats(), "companies": len(companies), "operators": len(operators), "enriched": len(to_enrich)}


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Profile seller, triage + enrich event attendees, write the graph.")
    ap.add_argument("--seller", required=True)
    ap.add_argument("--sheet", required=True, type=Path)
    ap.add_argument("--event", required=True)
    ap.add_argument("--deep", type=int, default=30)
    args = ap.parse_args()
    print(json.dumps(run(args.seller, args.sheet, args.event, args.deep), indent=2))
